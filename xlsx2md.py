#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
xlsx2md.py — Excel(.xlsx / .xlsm) の各シートを Markdown の表(md表)に変換する

使い方(詳細は README.md):
    python xlsx2md.py                      # ./input の全Excelを ./output へ変換
    python xlsx2md.py 構成図チェック観点一覧.xlsx   # ファイル指定
    python xlsx2md.py input -o output       # フォルダ指定

変換ルール(要点):
  - 先頭(または --main-sheet で指定)のシート → <元ファイル名>.md
  - それ以外のシート            → <元ファイル名>_<シート名>.md
  - Excelのテーブル(ListObject)があればその範囲を表として使う。無ければ見出し行を自動判定
  - セル内の改行は <br>、半角の縦棒 | は全角 ｜ に置き換える(md表を崩さないため)
  - 非表示の行・列は出力しない(--include-hidden で出力)
  - セル内に他シートの名前が出てきたら、そのシートのmdへのリンクに置き換える(--no-link-refs で無効)
  - 表以外の行(タイトル・元ファイル名・関連シートのリンク)は | で始めない
    (行数をツールで数える約束: 「| で始まる行」=表の行)

必要なもの: Python 3.9 以上、openpyxl(Anaconda には同梱。無ければ pip install openpyxl)
"""

from __future__ import annotations

import argparse
import datetime as _dt
import re
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import get_column_letter, range_boundaries
except ImportError:  # pragma: no cover
    print("openpyxl が見つかりません。 pip install openpyxl を実行してください。", file=sys.stderr)
    sys.exit(1)


# ---------------------------------------------------------------------------
# 文字列の整形
# ---------------------------------------------------------------------------

_INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')


def sanitize_filename(name: str) -> str:
    """シート名をファイル名に使える形にする(Windowsで使えない記号を _ に)。"""
    return _INVALID_FILENAME_CHARS.sub("_", name).strip() or "sheet"


def cell_to_text(value) -> str:
    """セルの値を md表の1セルに入る文字列にする。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        # 3.0 → 3 、それ以外は不要な末尾ゼロを落とす
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, _dt.datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.strftime("%Y/%m/%d")
        return value.strftime("%Y/%m/%d %H:%M")
    if isinstance(value, _dt.date):
        return value.strftime("%Y/%m/%d")
    text = str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("|", "｜")          # 半角の縦棒は表の区切りと衝突する
    text = text.strip()
    text = re.sub(r"\n+", "<br>", text)     # セル内改行
    text = re.sub(r"[ \t]+", " ", text)
    return text


# ---------------------------------------------------------------------------
# 表の範囲の決定
# ---------------------------------------------------------------------------

def find_table_range(ws) -> tuple[int, int, int, int, str]:
    """シートの表の範囲 (min_col, min_row, max_col, max_row, 由来) を返す。

    1. Excelのテーブル(ListObject)があれば、その範囲(最初のテーブル)
    2. 無ければ、先頭30行の中で「空でないセルが最も多い行」を見出し行とし、
       そこからシート末尾までを表とみなす
    """
    if ws.tables:
        table = next(iter(ws.tables.values()))
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        return min_col, min_row, max_col, max_row, f"テーブル '{table.name}' ({table.ref})"

    best_row, best_count = None, 0
    scan_to = min(ws.max_row, 30)
    for r in range(1, scan_to + 1):
        count = sum(1 for c in ws[r] if c.value not in (None, ""))
        if count > best_count:
            best_row, best_count = r, count
    if best_row is None:
        raise ValueError("表が見つかりません(空のシート)")

    # 見出し行の中で値がある最初と最後の列
    header_cells = [c for c in ws[best_row] if c.value not in (None, "")]
    min_col = header_cells[0].column
    max_col = header_cells[-1].column
    return min_col, best_row, max_col, ws.max_row, f"見出し行を自動判定 (行 {best_row})"


def merged_value_map(ws) -> dict[tuple[int, int], object]:
    """結合セルの左上の値を、結合範囲の全セルに広げるための辞書。"""
    mapping: dict[tuple[int, int], object] = {}
    for rng in ws.merged_cells.ranges:
        top_left = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                mapping[(r, c)] = top_left
    return mapping


# ---------------------------------------------------------------------------
# 1シート → md
# ---------------------------------------------------------------------------

def sheet_to_rows(ws, include_hidden: bool, drop_columns: set[str]):
    """シートを (見出しリスト, 行リスト, 説明) に変換する。"""
    min_col, min_row, max_col, max_row, origin = find_table_range(ws)
    merged = merged_value_map(ws)

    def value_at(r: int, c: int):
        if (r, c) in merged:
            return merged[(r, c)]
        return ws.cell(row=r, column=c).value

    def col_hidden(c: int) -> bool:
        dim = ws.column_dimensions.get(get_column_letter(c))
        return bool(dim and dim.hidden)

    def row_hidden(r: int) -> bool:
        dim = ws.row_dimensions.get(r)
        return bool(dim and dim.hidden)

    # 見出し
    columns = []
    for c in range(min_col, max_col + 1):
        header = cell_to_text(value_at(min_row, c))
        if not include_hidden and col_hidden(c):
            continue
        if header in drop_columns:
            continue
        columns.append((c, header or get_column_letter(c)))

    # データ行(空行は飛ばす)
    rows = []
    for r in range(min_row + 1, max_row + 1):
        if not include_hidden and row_hidden(r):
            continue
        values = [cell_to_text(value_at(r, c)) for c, _ in columns]
        if all(v == "" for v in values):
            continue
        rows.append(values)

    return [h for _, h in columns], rows, origin


def link_sheet_refs(text: str, sheet_links: dict[str, str], self_name: str) -> str:
    """セル内に他シートの名前があれば、そのmdへのリンクに置き換える。"""
    for name, filename in sheet_links.items():
        if name == self_name or not name:
            continue
        if name in text and f"]({filename})" not in text:
            text = text.replace(name, f"[{name}]({filename})")
    return text


def render_markdown(title: str, source_name: str, headers: list[str], rows: list[list[str]],
                    related: list[tuple[str, str]], sheet_links: dict[str, str],
                    link_refs: bool, converted_at: str) -> str:
    lines = [f"# {title}", ""]
    lines.append(f"元ファイル: {source_name} / シート: {title} / 変換日時: {converted_at}")
    if related:
        rel = " / ".join(f"[{n}]({f})" for n, f in related)
        lines.append(f"関連シート: {rel}")
    lines.append("")
    lines.append("| " + " | ".join(headers) + " |")
    lines.append("|" + "|".join("---" for _ in headers) + "|")
    for row in rows:
        cells = row
        if link_refs:
            cells = [link_sheet_refs(v, sheet_links, title) for v in row]
        lines.append("| " + " | ".join(cells) + " |")
    lines.append("")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 1ブック → 複数md
# ---------------------------------------------------------------------------

def convert_workbook(xlsx_path: Path, out_dir: Path, main_sheet: str | None,
                     include_hidden: bool, include_hidden_sheets: bool,
                     drop_columns: set[str], link_refs: bool) -> list[Path]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=False)
    base = xlsx_path.stem
    converted_at = _dt.datetime.now().strftime("%Y/%m/%d %H:%M")

    sheets = [ws for ws in wb.worksheets if include_hidden_sheets or ws.sheet_state == "visible"]
    if not sheets:
        raise ValueError(f"{xlsx_path.name}: 表示中のシートがありません")

    if main_sheet:
        names = [ws.title for ws in sheets]
        if main_sheet not in names:
            raise ValueError(f"{xlsx_path.name}: シート '{main_sheet}' がありません。あるシート: {names}")
        main = next(ws for ws in sheets if ws.title == main_sheet)
    else:
        main = sheets[0]

    # 出力ファイル名の対応表(シート名 → ファイル名)
    sheet_links: dict[str, str] = {}
    for ws in sheets:
        if ws is main:
            sheet_links[ws.title] = f"{base}.md"
        else:
            sheet_links[ws.title] = f"{base}_{sanitize_filename(ws.title)}.md"

    written: list[Path] = []
    for ws in sheets:
        try:
            headers, rows, origin = sheet_to_rows(ws, include_hidden, drop_columns)
        except ValueError as e:
            print(f"  - スキップ: シート '{ws.title}' ({e})")
            continue
        related = [(n, f) for n, f in sheet_links.items() if n != ws.title]
        md = render_markdown(ws.title, xlsx_path.name, headers, rows, related,
                             sheet_links, link_refs, converted_at)
        out_path = out_dir / sheet_links[ws.title]
        out_path.write_text(md, encoding="utf-8", newline="\n")
        written.append(out_path)
        print(f"  - {out_path.name}: 表の行数 {len(rows)} / 列数 {len(headers)} / {origin}")
    return written


def collect_inputs(target: Path) -> list[Path]:
    if target.is_file():
        return [target]
    files = sorted(p for p in target.iterdir()
                   if p.suffix.lower() in (".xlsx", ".xlsm") and not p.name.startswith("~$"))
    return files


def main(argv=None) -> int:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(
        description="Excel(.xlsx/.xlsm)の各シートをMarkdownの表に変換する",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("target", nargs="?", default=str(script_dir / "input"),
                        help="変換するExcelファイル、またはExcelが入ったフォルダ(既定: ./input)")
    parser.add_argument("-o", "--output", default=str(script_dir / "output"),
                        help="出力先フォルダ(既定: ./output)")
    parser.add_argument("--main-sheet", default=None,
                        help="<元ファイル名>.md にするシート名(既定: 先頭の表示中シート)")
    parser.add_argument("--drop-columns", default="",
                        help="出力しない列の見出し(カンマ区切り。例: '判定トリガー（削除予定）,メモ')")
    parser.add_argument("--include-hidden", action="store_true",
                        help="非表示の行・列も出力する(既定: 出力しない)")
    parser.add_argument("--include-hidden-sheets", action="store_true",
                        help="非表示のシートも変換する(既定: 変換しない)")
    parser.add_argument("--no-link-refs", action="store_true",
                        help="セル内の他シート名をリンクに置き換えない")
    args = parser.parse_args(argv)

    target = Path(args.target)
    out_dir = Path(args.output)
    if not target.exists():
        print(f"見つかりません: {target}", file=sys.stderr)
        return 1
    out_dir.mkdir(parents=True, exist_ok=True)

    files = collect_inputs(target)
    if not files:
        print(f"Excelファイルがありません: {target}", file=sys.stderr)
        return 1

    drop_columns = {c.strip() for c in args.drop_columns.split(",") if c.strip()}
    total = 0
    for xlsx in files:
        print(f"変換: {xlsx.name}")
        try:
            written = convert_workbook(xlsx, out_dir, args.main_sheet, args.include_hidden,
                                       args.include_hidden_sheets, drop_columns,
                                       not args.no_link_refs)
        except PermissionError:
            print(f"  ! 開けません(Excelで開いたままの可能性): {xlsx.name}", file=sys.stderr)
            continue
        except Exception as e:  # noqa: BLE001
            print(f"  ! 失敗: {xlsx.name}: {e}", file=sys.stderr)
            continue
        total += len(written)
    print(f"完了: {total} ファイルを {out_dir} に出力しました")
    return 0


if __name__ == "__main__":
    sys.exit(main())
